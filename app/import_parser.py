"""
Bring-your-own test cases: deterministic tabular parsing for .xlsx/.csv.

Tried FIRST for spreadsheet-shaped files, before ever calling an LLM - if
the file has a header row we can confidently match, parsing it this way is
free, instant, and can't misread a well-formed table the way an AI
structuring pass theoretically could. Returns None (not an empty list) when
headers can't be confidently matched, so the caller knows to fall back to
qa_engine.structure_existing_test_cases() instead of silently returning
nothing.
"""
import csv
import io
import re
from typing import Optional


def _normalize(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (header or "").lower())


# Each field maps to a set of normalized header aliases we'll recognize.
# Deliberately conservative - a near-miss falls through to the AI path
# rather than risk guessing the wrong column.
_FIELD_ALIASES = {
    "tc_id": {"tcid", "testcaseid", "testid", "caseid", "id"},
    "title": {"title", "name", "testcase", "testcasename", "testname", "summary"},
    "precondition": {"precondition", "preconditions", "prerequisite", "prerequisites", "setup"},
    "steps": {"steps", "teststeps", "stepstoreproduce", "procedure", "action", "actions"},
    "expected_result": {"expectedresult", "expected", "expectedoutcome", "result", "expectedbehavior"},
}


def _match_headers(headers: list) -> Optional[dict]:
    """Returns {field_name: column_index} if title and steps are both
    confidently found (the two fields a test case can't meaningfully be
    without); other fields are optional. None if that minimum isn't met."""
    normalized = [_normalize(h) for h in headers]
    mapping = {}
    for field, aliases in _FIELD_ALIASES.items():
        for idx, h in enumerate(normalized):
            if h in aliases:
                mapping[field] = idx
                break
    if "title" in mapping and "steps" in mapping:
        return mapping
    return None


def _rows_to_cases(rows: list, mapping: dict) -> list:
    cases = []
    for i, row in enumerate(rows):
        def get(field):
            idx = mapping.get(field)
            if idx is None or idx >= len(row):
                return ""
            val = row[idx]
            return "" if val is None else str(val).strip()

        title = get("title")
        steps = get("steps")
        if not title and not steps:
            continue  # skip genuinely blank rows
        cases.append({
            "tc_id": get("tc_id") or f"TC-{i+1:03d}",
            "title": title,
            "precondition": get("precondition"),
            "steps": steps,
            "expected_result": get("expected_result"),
            "notes": "",
        })
    return cases


def try_parse_tabular(filename: str, raw_bytes: bytes) -> Optional[list]:
    """Returns a list of test-case dicts if this looks like a well-headered
    spreadsheet of test cases, or None if the format/headers aren't
    recognizable (caller should fall back to AI structuring in that case)."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "csv":
        try:
            text = raw_bytes.decode("utf-8", errors="ignore")
            reader = csv.reader(io.StringIO(text))
            rows = list(reader)
        except Exception:
            return None
        if not rows:
            return None
        mapping = _match_headers(rows[0])
        if not mapping:
            return None
        return _rows_to_cases(rows[1:], mapping)

    if ext in ("xlsx", "xlsm"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw_bytes), data_only=True, read_only=True)
            ws = wb.worksheets[0]
            rows = list(ws.iter_rows(values_only=True))
        except Exception:
            return None
        if not rows:
            return None
        mapping = _match_headers([str(h) if h is not None else "" for h in rows[0]])
        if not mapping:
            return None
        return _rows_to_cases(rows[1:], mapping)

    return None
